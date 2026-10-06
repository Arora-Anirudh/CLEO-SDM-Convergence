#!/usr/bin/env python3
"""Build the complete 2-D constthermo2d fixed-grid experiment report."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/reports/Two_Dimensional_CLEO_Condensation_Collision_Convergence_Report_2026-09-28_layout_fixed.docx"
ANALYSIS = ROOT / "results/analysis"


def set_cell_fill(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    node = OxmlElement("w:keepNext")
    p_pr.append(node)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))


def math_run(text: str):
    run = OxmlElement("m:r")
    props = OxmlElement("m:rPr")
    font = OxmlElement("m:ctrlPr")
    props.append(font)
    run.append(props)
    node = OxmlElement("m:t")
    node.text = text
    run.append(node)
    return run


def math_sub(base: str, sub: str):
    node = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    e.append(math_run(base))
    s = OxmlElement("m:sub")
    s.append(math_run(sub))
    node.extend((e, s))
    return node


def math_sup(base: str, sup: str):
    node = OxmlElement("m:sSup")
    e = OxmlElement("m:e")
    e.append(math_run(base))
    s = OxmlElement("m:sup")
    s.append(math_run(sup))
    node.extend((e, s))
    return node


def math_subsup(base: str, sub: str, sup: str):
    node = OxmlElement("m:sSubSup")
    e = OxmlElement("m:e")
    e.append(math_run(base))
    ssub = OxmlElement("m:sub")
    ssub.append(math_run(sub))
    ssup = OxmlElement("m:sup")
    ssup.append(math_run(sup))
    node.extend((e, ssub, ssup))
    return node


def math_fraction(numerator, denominator):
    frac = OxmlElement("m:f")
    num = OxmlElement("m:num")
    den = OxmlElement("m:den")
    for item in numerator:
        num.append(item)
    for item in denominator:
        den.append(item)
    frac.extend((num, den))
    return frac


def add_equation(doc: Document, parts) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    for part in parts:
        omath.append(part)
    omath_para.append(omath)
    paragraph._p.append(omath_para)


def setup_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, space_before, space_after in (
        ("Title", 24, "000000", 0, 12),
        ("Subtitle", 13, "3A4C62", 0, 16),
        ("Heading 1", 16, "000000", 16, 7),
        ("Heading 2", 12.5, "000000", 12, 5),
        ("Heading 3", 11, "000000", 9, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Heading 3" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(space_after)
        style.paragraph_format.keep_with_next = True

    if "Figure Caption" not in [s.name for s in doc.styles]:
        caption = doc.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = doc.styles["Figure Caption"]
    caption.base_style = doc.styles["Normal"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor(45, 55, 72)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.line_spacing = 1.0

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.text = "CLEO constthermo2d fixed grid convergence study"
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(85, 98, 112)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("28 September 2026  |  Page ")
    add_page_number(footer)
    for run in footer.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(85, 98, 112)
    return doc


def add_title(doc: Document) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Two Dimensional CLEO Condensation Collision Convergence Experiment")
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Fixed grid normal sampling 20 member resolution study")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(22)
    run = p.add_run("Complete methods results interpretation and reproducibility report\n28 September 2026")
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(80, 94, 110)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    shade_paragraph(p, "EEF3F7")
    r = p.add_run("Main conclusion. ")
    r.bold = True
    p.add_run(
        "For the audited 120 minute 2-D constthermo2d configuration, "
        "1,024 superdroplets per initially populated grid box is the smallest tested "
        "resolution that satisfies the stated distribution, moment, speed, and terminal-speed "
        "transport checks when compared with the 8,192-SD-per-cell internal comparator and "
        "two adjacent resolution doublings. It corresponds to 122,880 initial SDs across the domain."
    )
    p = doc.add_paragraph()
    shade_paragraph(p, "FFF5E6")
    r = p.add_run("Interpretation boundary. ")
    r.bold = True
    p.add_run(
        "This is an operational resolution selection for this particular normal-sampling, "
        "fixed-grid, 120 minute trajectory. The model has null boundaries, so neither the "
        "terminal-speed transport proxy nor the fast-drop timing metric is surface rainfall or rainfall onset."
    )
    doc.add_page_break()


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        cell.text = text
        set_cell_fill(cell, "1F4E79")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9)
        if widths:
            cell.width = Inches(widths[index])
    for r_index, row_values in enumerate(rows):
        row = table.add_row()
        for index, text in enumerate(row_values):
            cell = row.cells[index]
            cell.text = text
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_index % 2:
                set_cell_fill(cell, "F4F7FA")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index and len(text) < 18 else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.size = Pt(8.8)
            if widths:
                cell.width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def add_figure(doc: Document, image: Path, number: int, title: str, explanation: str, width: float = 6.72) -> None:
    if not image.is_file():
        raise FileNotFoundError(image)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(image), width=Inches(width))
    cap = doc.add_paragraph(style="Figure Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    lead = cap.add_run(f"Figure {number}. {title}. ")
    lead.bold = True
    cap.add_run(explanation)


def add_figure_page(doc: Document, *args, **kwargs) -> None:
    doc.add_page_break()
    add_figure(doc, *args, **kwargs)


def add_report() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = setup_doc()
    add_title(doc)

    doc.add_heading("1 Purpose and decision", level=1)
    doc.add_paragraph(
        "This report consolidates the full 2-D CLEO constthermo2d fixed-grid "
        "condensation collision experiment. The scientific question is practical but specific: "
        "how many superdroplets per initially populated grid box are required before the "
        "ensemble-mean evolution of the droplet distribution and its physically meaningful "
        "diagnostics is stable for this simulation? The study begins from the upstream CLEO "
        "example and changes the sampling resolution only."
    )
    doc.add_paragraph(
        "The study establishes an operational, configuration-specific answer rather than a universal "
        "SDM resolution. The reference is the independent 20-member 8,192-SD-per-cell ensemble, "
        "not an analytic solution or an infinitely resolved calculation. The decision is therefore "
        "a reproducible threshold for follow-on experiments with the same domain, thermodynamics, "
        "kernel, duration, initialization family, and diagnostic suite."
    )

    doc.add_heading("2 Configuration inherited from the CLEO example", level=1)
    doc.add_paragraph(
        "The experiment uses CLEO v0.69.1 at upstream commit 9224585d784e9371926db9d0a23891a4b0241e24. "
        "The project-owned baseline preserves the upstream constthermo2d grid and physics; it does not "
        "change the domain geometry. The only production parameter deliberately varied is the initial "
        "number of superdroplets per populated grid box, denoted Ncell."
    )
    add_table(doc,
        ["Component", "Fixed setting", "Why it matters"],
        [
            ["Domain and grid", "1500 m x 1500 m x 20 m; 20 x 20 x-z boxes", "Each map cell is 75 m x 75 m x 20 m. The blocky spatial maps therefore show native model resolution."],
            ["Initial cloud layer", "0 to 500 m; 120 populated boxes", "The domain begins with 120 Ncell superdroplets. At Ncell=1024 this is 122,880 initial SDs."],
            ["Initial number concentration", "1.0 x 10^7 m^-3", "The prescribed physical droplet population represented by the SD multiplicities."],
            ["Radius proposal", "log-uniform in log10(r), nominally 3 nm to 3 micrometres", "The proposal used to sample SD radii; multiplicity recovers the prescribed bimodal physical distribution."],
            ["Thermodynamics", "prescribed dry hydrostatic adiabatic profile; saturation ratio 0.99 below and 1.0025 above 750 m", "Condensation and evaporation respond to this fixed environment; thermodynamic feedback is disabled."],
            ["Microphysics and motion", "condensation evaporation; Long collision coalescence; Rogers Gunn Kinzer terminal speed; Cartesian motion", "Connects DSD evolution to physical drop growth, transport, and collision outcomes."],
            ["Output cadence", "0 to 7200 s, saved every 120 s", "61 stored times, so timing diagnostics have a 2-minute sampling precision."],
            ["Boundary condition", "NullBoundaryConditions", "Drops and mass do not leave through a lower boundary. Spatial movement and virtual flux are not surface precipitation."],
        ], [1.25, 2.25, 3.05])

    doc.add_heading("3 Sampling resolution and ensemble design", level=1)
    doc.add_paragraph(
        "The normal-sampling production ladder uses 20 independent seeded members at every completed "
        "resolution. The base cohort covers Ncell = 2, 4, 8, ..., 4,096. The 8,192-SD-per-cell extension "
        "was generated as a separate fresh 20-member high-resolution cohort. Pilot, smoke, gate, and thread "
        "benchmark runs are retained as execution evidence but are not pooled into these ensembles."
    )
    add_table(doc,
        ["Resolution range", "Members per level", "Initial domain SD count", "Use in analysis"],
        [
            ["2 through 4,096", "20", "120 Ncell", "Primary resolution ladder and adjacent-doubling comparisons."],
            ["8,192", "20", "983,040", "Independent high-resolution internal comparator."],
            ["Ncell=1,024", "20", "122,880", "Smallest level that passes the stated original and physical convergence checks."],
        ], [1.45, 1.05, 1.65, 2.4])
    doc.add_paragraph(
        "The 8,192 cohort required a narrow support correction before model integration because the original "
        "3 nm to 3 micrometre proposal occasionally assigned zero integer multiplicity in the extreme tails. "
        "Only its proposal support was changed to 4 nm to 2 micrometres; the grid, physical target, kernel, "
        "thermodynamics, cadence, and seed protocol were retained. The analytical support audit estimates that "
        "the omitted probability is 0.000052% of initial number and 0.00205% of initial liquid mass. The high "
        "resolution cohort is therefore called a support-corrected internal comparator, not exact numerical truth."
    )

    doc.add_heading("4 How the CLEO trajectories were run and audited", level=1)
    doc.add_paragraph(
        "The executable uses a single MPI rank with the Kokkos Threads backend. Multi-rank MPI is not supported "
        "by this upstream example because no domain decomposition or exchange of moving superdroplets is implemented. "
        "The production shape was one Slurm task with 16 Kokkos Threads, 2 GiB memory, and at most 20 independently "
        "running array members. The 16-thread choice was cost-efficient in the Ncell=512 thread check; it is not a claim "
        "that it is universally fastest at every resolution."
    )
    add_table(doc,
        ["Item", "Location or command", "Audit role"],
        [
            ["Frozen configuration", "config/constthermo2d_fixedgrid_baseline_v1.yaml", "Preserves the upstream grid, physics, cadence, and fixed settings."],
            ["Production member list", "config/constthermo2d_fixedgrid_production_v1.csv", "Records the 20 seeds and expected initial totals at each resolution."],
            ["Preparation", "scripts/levante/prepare_constthermo2d_fixedgrid_production_configs_v1.sbatch", "Materialises immutable resolution YAML files and SHA-256 receipts before arrays run."],
            ["CLEO member wrapper", "scripts/levante/run_constthermo2d_fixedgrid_production_member_v1.sbatch", "Runs one fresh seeded case and invokes the trajectory integrity audit."],
            ["Analysis wrapper", "scripts/levante/analyze_constthermo2d_fixedgrid_production_array_v1.sbatch", "Streams accepted raw output to compact audited member summaries."],
            ["Raw outputs and logs", "/scratch/m/m301324/SDM/constthermo2d_fixedgrid_baseline_v1/", "Contains case output Zarr stores, integrity receipts, and array stdout and stderr logs."],
        ], [1.35, 3.05, 2.15])
    doc.add_paragraph(
        "The completed base ladder contains 240 of 240 passed trajectory-integrity receipts. Its observed allocated "
        "cost was 123.502 CPU-hours, within the pre-submission expectation of 115 to 130 CPU-hours. A preparation race "
        "was observed early in the Ncell=4 wave; it occurred before CLEO execution and was repaired by pre-materialising "
        "all immutable configurations. Failed control-plane logs were preserved, and only absent members were rerun."
    )

    doc.add_heading("5 Diagnostic definitions", level=1)
    doc.add_paragraph(
        "Every diagnostic is multiplicity-weighted. A superdroplet is a computational representative, while its "
        "multiplicity xi is the number of physical droplets it represents. Consequently, the following moments are "
        "physical concentration estimates rather than unweighted counts of simulated particles.")
    add_equation(doc, [
        math_sub("λ", "k"), math_run("(t) = "),
        math_fraction([math_run("1")], [math_run("V")]),
        math_sub("Σ", "i"), math_sub("ξ", "i"), math_subsup("r", "i", "k"), math_run("(t)"),
    ])
    doc.add_paragraph(
        "Here V is domain volume, ri is wet radius, and xi is multiplicity. Lambda0 is number concentration; lambda2 "
        "weights geometric surface area; lambda3 weights liquid volume; and lambda6 strongly emphasizes rare large drops."
    )
    add_equation(doc, [
        math_run("A = 4π"), math_sub("λ", "2"), math_run(",     L = "), math_sub("ρ", "w"), math_run(" (4π/3) "), math_sub("λ", "3"), math_run(",     Z = 64 "), math_sub("λ", "6"), math_run(" (r in mm)")
    ])
    doc.add_paragraph(
        "A is geometric droplet surface-area density, a geometric condensation-relevant quantity rather than a complete "
        "radiative-transfer extinction calculation. L is liquid-equivalent water content. Z is a Rayleigh reflectivity-factor "
        "proxy, reported in dBZ for interpretability; it is not a radar forward simulation."
    )
    add_equation(doc, [
        math_run("P(t) = "), math_fraction([math_run("1")], [math_run("V")]),
        math_sub("Σ", "i"), math_sub("ξ", "i"), math_sub("m", "i"), math_sub("v", "t"), math_run("("), math_sub("r", "i"), math_run(")")
    ])
    doc.add_paragraph(
        "P is the terminal-speed mass-transport proxy in g m-2 s-1. It asks how much represented liquid mass would move "
        "downward per unit area and time if each drop travelled at terminal speed. Because this configuration has null boundaries, "
        "P is not measured surface mass flux."
    )
    add_equation(doc, [
        math_run("δ(t) = "), math_fraction([math_run("200 |A(t) − B(t)|")], [math_run("|A(t)| + |B(t)|")]),
        math_run(";     TV(t) = 50 "), math_sub("Σ", "b"), math_run("|"), math_sub("p", "b"), math_run("(t) − "), math_sub("q", "b"), math_run("(t)|")
    ])
    doc.add_paragraph(
        "The scalar comparison uses the symmetric percent difference delta. The distribution statistic TV compares the "
        "500-bin mass DSD after each DSD is normalised to unit liquid mass. For each pair of resolutions, 20 members on both "
        "sides are independently resampled with replacement 2,000 times; the one-sided 95% upper bound of the maximum "
        "post-initial difference is compared with the stated margin."
    )

    doc.add_heading("6 Operational convergence rule", level=1)
    doc.add_paragraph(
        "A candidate must pass against the 8,192 internal comparator and across both adjacent transitions N to 2N and 2N to 4N. "
        "The core 5% margin applies to 500-bin mass-DSD total variation, lambda0, lambda2, lambda3, and liquid-mass-weighted "
        "terminal speed. A 10% tail margin applies to lambda6 and mass-weighted r90. The 250- and 1,000-bin mass-DSD calculations "
        "are representation-sensitivity checks rather than separate decision metrics."
    )
    add_table(doc,
        ["Candidate", "Comparison with 8,192", "N to 2N", "2N to 4N", "Decision"],
        [
            ["2 to 256", "At least one required metric exceeds its margin", "Fails", "Fails", "Reject"],
            ["512", "Passes comparator check", "5.11% mass-DSD bound exceeds 5%", "Passes", "Reject"],
            ["1,024", "Largest ratio 0.60", "Largest ratio 0.75", "Largest ratio 0.59", "Select"],
            ["2,048", "Passes", "Passes", "Passes", "Higher-resolution confirmation"],
        ], [1.0, 1.55, 1.35, 1.35, 1.35])
    doc.add_paragraph(
        "The later physical corroboration is deliberately independent in meaning, although it uses the same underlying trajectories. "
        "It rejects 512 because the 512-to-1,024 transport upper bound is 6.46% against the 5% margin, while 1,024-to-8,192 "
        "passes with a 3.53% transport bound. This protects the selected resolution from being justified only by conventional moments."
    )

    doc.add_heading("7 Complete figure record and interpretation", level=1)
    doc.add_paragraph(
        "The following pages contain the canonical, latest-version figures for this experiment. Earlier duplicate versions were "
        "superseded during visual and methodological checks and are not repeated. All curves labelled as ensemble means average "
        "the 20 independent members available at that resolution."
    )

    figures = [
        (ANALYSIS / "constthermo2d_fixedgrid_with8192_descriptive_v4/01_ensemble_moment_and_physical_evolution.png", "Ensemble mean evolution across the full resolution ladder", "This establishes the broad resolution trend before any decision rule is applied. Low-resolution curves visibly separate during the collision and transport transition, whereas the high-resolution curves gather together. It is descriptive evidence; visual overlap alone is not the convergence test."),
        (ANALYSIS / "constthermo2d_fixedgrid_with8192_descriptive_v4/02_mass_dsd_evolution_representative_resolutions.png", "Mass distribution evolution at representative resolutions", "The mass DSD weights each radius bin by represented liquid mass, so it makes the transfer of liquid into large drops explicit. The comparison shows why distribution-based evidence is needed alongside moments: similar bulk mass can coexist with different allocation of mass across radius."),
        (ANALYSIS / "constthermo2d_fixedgrid_with8192_descriptive_v4/03_descriptive_full20_reference_ladder.png", "Descriptive agreement with the internal 8,192 comparator", "Each point summarizes a full-20 comparison against the high-resolution ensemble. Differences fall with increasing Ncell, with the tail-sensitive quantities decreasing more slowly because rare large drops dominate them."),
        (ANALYSIS / "constthermo2d_fixedgrid_with8192_descriptive_v4/04_member_violin_diagnostics_at_multiple_times.png", "Member level diagnostic distributions at several times", "The violins expose the collision-realization spread hidden by an ensemble mean. Narrower distributions at higher Ncell indicate that both sampling resolution and finite-ensemble uncertainty improve together, although this figure itself is not a bootstrap confidence interval."),
        (ANALYSIS / "constthermo2d_fixedgrid_with8192_descriptive_v4/05_resolution_ensemble_precision_surface.png", "Resolution and ensemble precision surface", "Moving upward changes superdroplet resolution; moving right averages more independently random members. The colours therefore separate two sources of imprecision: discretisation and finite-ensemble sampling. This is a descriptive resampling surface, not an extrapolated fit beyond the 20 members actually run."),
        (ANALYSIS / "constthermo2d_fixedgrid_with8192_bootstrap_v4/06_high_resolution_bootstrap_bounds.png", "High resolution bootstrap bounds", "Observed all-time differences and their one-sided 95% member-bootstrap upper bounds are displayed separately. The bound is deliberately conservative because it allows for re-sampling variation among the 20 members."),
        (ANALYSIS / "constthermo2d_fixedgrid_with8192_bootstrap_v4/07_adjacent_doubling_dsd_bin_sensitivity.png", "Adjacent doubling and mass DSD bin sensitivity", "The three fixed radius meshes test whether the conclusion depends on distribution representation. The 500-bin DSD is the decision representation; the 250- and 1,000-bin curves provide a useful check that bin choice does not drive the observed improvement with resolution."),
        (ANALYSIS / "constthermo2d_fixedgrid_operational_selection_v6/01_operational_convergence_gate_matrix.png", "Original multi metric operational convergence gate", "This is the primary selection figure. A candidate must remain below its applicable margin in all required comparisons. Ncell=512 is close but fails its first adjacent doubling, whereas 1,024 is the first level to pass the full gate."),
        (ANALYSIS / "constthermo2d_fixedgrid_physical_diagnostics_v4/01_all_moments_and_mass_evolution.png", "All moments and represented droplet mass", "The panel suite follows lambda0 through lambda6 and total mass. Increasing resolution primarily improves the stability of the collision-generated tail; lambda6 is the clearest example because each large drop is amplified by r to the sixth power."),
        (ANALYSIS / "constthermo2d_fixedgrid_physical_diagnostics_v4/02_radius_transport_and_population_evolution.png", "Radii terminal speed lower layer mass and active superdroplets", "Mass-weighted r50, r90, and r99 describe where liquid mass sits in the spectrum. The terminal-speed panel converts that spectrum into a dynamically meaningful speed summary. Mass below 150 m is an in-domain location diagnostic, while active SD count is computational bookkeeping rather than physical number concentration."),
        (ANALYSIS / "constthermo2d_fixedgrid_physical_diagnostics_v4/03_number_dsd_evolution.png", "Number DSD evolution", "This DSD counts represented droplets per logarithmic radius interval. The jagged two-SD spectrum is a direct illustration of insufficient sampling. High-resolution shapes are much more stable, especially through the transition to larger-drop populations."),
        (ANALYSIS / "constthermo2d_fixedgrid_physical_diagnostics_v4/04_total-droplet-mass_dsd_evolution.png", "Mass DSD evolution", "Mass weighting makes the physically relevant redistribution into larger drops visible. This is the distribution used in the primary gate because collision coalescence changes how conserved liquid mass is apportioned between sizes."),
        (ANALYSIS / "constthermo2d_fixedgrid_physical_diagnostics_v4/05_physical_moment_proxies_evolution.png", "Surface area liquid water and reflectivity proxies", "The r2-derived area proxy connects to available droplet surface, r3-derived liquid water connects to represented volume, and r6-derived Z exposes the large-drop tail. The stronger separation in Z is physically expected rather than a plotting problem."),
        (ANALYSIS / "constthermo2d_fixedgrid_physical_diagnostics_v4/06_virtual_flux_and_fastdrop_evolution.png", "Terminal speed transport and fast drop liquid fractions", "The transport proxy is a mass-weighted terminal-speed flux diagnostic, while the right-hand panels show the fraction of liquid in speed-defined classes. At the selected resolution and the 8,192 comparator, the class histories are close; thresholds above about 1 m s-1 are not informative in this 120-minute output."),
        (ANALYSIS / "constthermo2d_fixedgrid_physical_diagnostics_v4/07_terminal_speed_mass_spectrum.png", "Liquid mass spectrum in terminal speed coordinates", "This remaps the mass DSD from radius to the Rogers Gunn Kinzer terminal-speed relation. It is not an independently simulated velocity distribution. It shows that the relevant occupied fast-drop tail is predominantly below approximately 1 m s-1."),
        (ANALYSIS / "constthermo2d_fixedgrid_physical_diagnostics_v4/08_spatial_total_droplet_mass_fields.png", "Spatial total droplet mass fields", "Rows compare Ncell=2, 1,024, and 8,192 ensemble means at 0, 60, and 120 min. Each coloured square is one native 75 m by 75 m cell. The 1,024 and 8,192 patterns agree far more closely than the two-SD field, but the panels diagnose in-domain redistribution, not rain-out."),
        (ANALYSIS / "constthermo2d_fixedgrid_physical_diagnostics_v4/09_physical_diagnostics_reference_ladder.png", "Physical diagnostics reference ladder", "Dots are observed all-time differences against 8,192; bars are one-sided 95% member-bootstrap upper bounds; dashed lines are the operational margins. The 512-to-1,024 transport failure and the 1,024-to-8,192 pass are the key independent physical corroboration."),
        (ANALYSIS / "constthermo2d_fixedgrid_physical_diagnostics_v4/10_fastdrop_development_time.png", "Fast drop development time", "The metric is the first stored time when at least 1% of liquid mass has terminal speed of at least 1 m s-1. It is a useful emergence indicator with 2-minute temporal resolution. It must not be called rainfall onset because the configuration has no outflow boundary or surface collection."),
        (ANALYSIS / "constthermo2d_fixedgrid_size_class_visuals_v6/03_mass_dsd_time_radius_Ncell1024.png", "Time radius mass DSD at the selected resolution", "The colour field gives a compact view of where liquid mass occupies radius-time space. The overlaid cloud, drizzle-sized, and rain-sized size-class thresholds refer only to drop size: they do not imply drops have reached a surface."),
        (ANALYSIS / "constthermo2d_fixedgrid_size_class_visuals_v6/04_size_class_mass_fractions.png", "Mass fractions by cloud drizzle and rain sized classes", "This follows liquid mass as it migrates from the initial small-drop population to increasingly large size classes. The figure is complementary to the mass DSD: it is easy to discuss, while the DSD retains the full spectral information."),
        (ANALYSIS / "constthermo2d_fixedgrid_size_class_visuals_v6/05_native_grid_cell_mass_fields_Ncell1024.png", "Native grid cell mass fields at Ncell 1024", "This explicitly labels the native grid rather than visually smoothing it. The field is coarse because the physical model grid is 20 by 20, not because the image is low resolution."),
        (ANALYSIS / "constthermo2d_fixedgrid_spatial_mass_animation_v2/spatial_total_droplet_mass_evolution_first_frame.png", "Spatial liquid mass animation representative first frame", "The accompanying animation contains all 61 saved 120-second outputs, not interpolated images. It compares full-20 ensemble-mean fields at Ncell=2, 1,024, and 8,192 on one fixed logarithmic mass-concentration scale. The MP4 and GIF are retained beside this report's source figures."),
    ]
    for number, (image, title, explanation) in enumerate(figures, start=1):
        if number == 1:
            add_figure(doc, image, number, title, explanation)
        else:
            add_figure_page(doc, image, number, title, explanation)

    doc.add_page_break()
    doc.add_heading("8 Final selection and scientific interpretation", level=1)
    doc.add_paragraph(
        "The complete evidence supports Ncell=1,024 as the smallest tested operational resolution for the present experiment. "
        "It is selected only after passing the original multi-metric mass-DSD and moment gate, two adjacent resolution doublings, "
        "and the later physical transport audit. The immediate lower candidate, Ncell=512, is scientifically close but not selected: "
        "its adjacent mass-DSD bound is 5.11% against the 5% core margin, and its terminal-speed transport bound is 6.46% against "
        "the same 5% physical margin."
    )
    doc.add_paragraph(
        "The result says that an ensemble mean built from 20 independently seeded 1,024-SD-per-populated-cell runs reproduces the "
        "tested high-resolution internal comparator within the stated accuracy objectives. It does not say that every individual "
        "1,024-SD realization is identical, that 1,024 is universal for a different grid or thermodynamic setting, or that the "
        "simulation has been verified against an external numerical truth."
    )
    doc.add_heading("9 Limits that should accompany every presentation of the result", level=1)
    for text in (
        "The 8,192 ensemble is an internal support-corrected comparator, not an analytic or independently converged reference solution.",
        "The resolution rule and its 5% and 10% margins are project-level operational choices applied retrospectively to the completed ensemble, not universal physical tolerances.",
        "The 20-member ensemble quantifies one finite ensemble size. The precision surfaces are descriptive and do not prove extrapolated behaviour at larger ensemble sizes.",
        "Null boundaries prevent a surface rain-rate, rainfall onset, or surface flux conclusion. Fast-drop emergence and terminal-speed mass transport are physically motivated in-domain proxies only.",
        "The spatial field uses the example's native 75 m by 75 m cells. It should not be cosmetically smoothed and interpreted as subgrid structure.",
        "This report concerns normal sampling. Its selected resolution must not be transferred automatically to a future alpha-sampling experiment without a separate convergence audit.",
    ):
        add_bullet(doc, text)

    doc.add_heading("10 Reproducibility map", level=1)
    doc.add_paragraph(
        "The following artefacts are the authoritative provenance map for reproducing or extending this work. Raw Zarr stores remain "
        "authoritative; the local cache is a compact, audited reconstruction used for the analysis figures."
    )
    add_table(doc,
        ["Purpose", "Canonical artefact"],
        [
            ["Experiment design", "docs/experiments/constthermo2d_fixedgrid_convergence_design_v1.md"],
            ["Production execution record", "docs/runs/constthermo2d_fixedgrid_production_v1_submission_plan.md"],
            ["High-resolution extension record", "docs/runs/constthermo2d_fixedgrid_n8192_extension_plan_v1.md"],
            ["Compact member cache", "results/cache/constthermo2d_fixedgrid_production_v1/"],
            ["Core production aggregation", "scripts/analyze_constthermo2d_fixedgrid_production_v1.py"],
            ["Operational gate", "scripts/select_constthermo2d_fixedgrid_convergence_v1.py"],
            ["Physical diagnostics", "scripts/analyze_constthermo2d_fixedgrid_physical_v1.py"],
            ["DSD size class visualisations", "scripts/animate_constthermo2d_dsd_size_classes_v1.py"],
            ["Spatial animation export and renderer", "scripts/export_constthermo2d_spatial_mass_animation_inputs_v1.py and scripts/animate_constthermo2d_spatial_mass_v1.py"],
            ["Final canonical physical figure set", "results/analysis/constthermo2d_fixedgrid_physical_diagnostics_v4/"],
        ], [2.0, 4.8])
    doc.add_paragraph(
        "To recreate the native spatial animation, the two extraction inputs are the raw massmom1 observer arrays for all 20 members "
        "at Ncell=2, 1,024, and 8,192. They are averaged before animation, preserving each of the 61 model outputs. The stored MP4 is "
        "results/analysis/constthermo2d_fixedgrid_spatial_mass_animation_v2/spatial_total_droplet_mass_evolution_Ncell2_1024_8192.mp4."
    )

    doc.add_heading("Appendix A Figure inventory", level=1)
    doc.add_paragraph(
        "This report contains every latest canonical static figure created for the completed normal-sampling 2-D fixed-grid experiment: "
        "the five descriptive ensemble products, two high-resolution bootstrap products, the operational gate, ten physical diagnostics, "
        "three size-class products, and the representative frame for the full spatial animation. Earlier superseded figures and failed "
        "rendering attempts are intentionally not duplicated."
    )
    doc.add_heading("Appendix B Terminology", level=1)
    add_table(doc,
        ["Term", "Meaning in this report"],
        [
            ["Ncell", "Initial number of superdroplets in each of the 120 populated grid boxes."],
            ["SD", "Computational superdroplet. Its multiplicity gives the represented physical droplet count."],
            ["Mass DSD", "Liquid mass per logarithmic wet-radius interval, divided by domain volume."],
            ["Number DSD", "Represented physical droplet number per logarithmic wet-radius interval, divided by domain volume."],
            ["Internal comparator", "The independent full-20 Ncell=8,192 ensemble used as the practical high-resolution comparison, not numerical truth."],
            ["Virtual flux", "Mass-weighted terminal-speed transport proxy within the model, not flux through a domain boundary."],
            ["Fast drop", "A drop whose terminal speed passes the stated threshold; it is not automatically a precipitating surface raindrop in this setup."],
        ], [1.55, 5.25])

    doc.core_properties.title = "Two Dimensional CLEO Condensation Collision Convergence Experiment"
    doc.core_properties.subject = "Complete fixed-grid normal-sampling convergence and physical diagnostic report"
    doc.core_properties.author = "Anirudh Arora"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    add_report()
